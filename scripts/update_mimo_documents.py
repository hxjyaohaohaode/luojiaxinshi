"""Produce new MiMo editions from the user's documents and normalized template."""
from pathlib import Path
import copy
import hashlib
import io
import json
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Twips

sys.path.insert(0,r'C:\Users\LENOVO\.codex\document-templates\user-standard')
import document_style as style

ROOT=Path(__file__).resolve().parents[2]
FOLDER=ROOT/'项目内容设计稿'
MODELS=[
    ('mimo-v2.6-pro','通知与附件理解、结构化提取、截图理解、问答、计划解释及智能体角色推理','统一文本与视觉主模型；事实和日期仍需原文校验。'),
    ('mimo-v2.5-asr','口述问题、录音转文字','转写后先回填，由本人核对人名、数字与期限后发送。'),
    ('mimo-v2.5-tts','默认通知、待办和回答播报','使用小米预置音色，按已核对文字合成。'),
    ('mimo-v2.5-tts-voicedesign','按文字描述设计小信使音色','描述与播报正文分开传入，关闭正文自动润色。'),
    ('mimo-v2.5-tts-voiceclone','本人或获授权声音的播报','提供参考录音并确认声音所有者授权；每次请求携带参考音频。'),
]
REPLACEMENTS={
    '首版采用单一模型和固定步骤。模型只产草稿，无权建项、访问其他账号或执行外部操作；通知中的操作指令视为待处理内容。服务端校验本人确认请求后才写入并防止重复提交。':
    '首版采用小米 mimo-v2.6-pro 作为统一文本与视觉主模型，配合固定处理步骤；语音识别和三种播报场景分别调用四个小米语音模型。主模型只产草稿，无权直接建项、访问其他账号或执行外部操作。服务端校验本人确认请求后才写入并防止重复提交。',
    '中；成熟解析组件配合一种获准使用的模型，逐文件校验后联合处理。':
    '中；成熟解析组件配合 mimo-v2.6-pro，逐文件校验后联合处理。',
    '视觉识别待资源与样本验证后增加；人工补文字单独计入成本和失败统计。':
    '截图由 mimo-v2.6-pro 理解，重要文字由本人核对；扫描 PDF 的整页识别流程须单独验证，人工补文字计入成本和失败统计。',
    '协调智能体组织任务，专业智能体可互发请求，进行澄清、约束协商、复核与交接。消息关联具体任务、当前事实、权限、时限及返回结果；协调器控制参与角色、往返次数和运行预算，避免重复讨论与失控执行。':
    '协调、通知分析、个人匹配、计划解释和复核等智能体角色统一由小米 mimo-v2.6-pro 驱动，通过不同任务提示及独立上下文组织澄清、约束协商与交接。消息关联具体任务、事实、权限、时限及结果；协调器限制往返次数和调用预算。语音能力按下表路由，不要求每次任务同时调用全部模型。',
    '文件与图片解析服务读取 PDF、Word、表格及截图；模型服务完成理解、结构化抽取和协作；检索、规则与排程服务负责关联、校验及约束计算。后端配置私有文件存储、数据库、持久任务队列、消息推送与运行监测，统一管理凭证和调用费用。':
    'PDF、Word、表格等可提取文字由本地解析组件读取，截图由 mimo-v2.6-pro 理解；全部语义分析、结构化抽取及智能体推理由同一小米主模型承担。语音输入使用 mimo-v2.5-asr，普通播报、音色设计和声音复刻分别使用对应 MiMo TTS 模型。检索、规则、权限、日期及排程约束仍由确定性程序处理；私有文件、数据库、队列与推送渠道沿用原架构，统一记录调用结果和费用。',
    '有模型 API 时，采用网页应用、服务端、数据库和私有文件目录，密钥仅存服务端。赛事平台实际具备模型调用、存储及访问控制时可复用；仅有聊天页面不等于具备这些能力。':
    '模型服务固定对接小米 MiMo 按量 API，采用网页应用、服务端、数据库和私有文件目录，MIMO_API_KEY 仅存服务端。模型调用均经过本项目后端，浏览器不持有密钥。赛事平台只有实际具备所需模型接口、存储及访问控制时才复用。',
    '模型超时、额度不足、结构错误或内容超限时，保留材料、显示失败，允许有限重试并记录调用和费用，不假报成功。模型资源未落实时，可开发页面与手工流程，AI 功能待验证。':
    '小米接口超时、额度不足、结构错误或内容超限时保留材料并显示失败，由用户手动重试可能计费的请求，不回退其他 AI 供应商。当前已预置五个模型及接口适配；API Key、账号权限和真实效果仍待逐模型联调。',
}


def finish_layout(doc):
    for table in doc.tables:
        widths=[c.width for c in table.columns]
        total=sum(w.twips for w in widths)
        width=table._tbl.tblPr.find(qn('w:tblW'))
        width.set(qn('w:type'),'dxa');width.set(qn('w:w'),str(total))
        is_models=table.cell(0,0).text=='模型'
        for ri,row in enumerate(table.rows):
            for ci,cell in enumerate(row.cells):
                cell.width=widths[ci]
                for p in cell.paragraphs:p.paragraph_format.keep_with_next=(ri==0 or is_models and ri<len(table.rows)-1)
    for i,p in enumerate(doc.paragraphs):
        if p.text=='小米模型分工':
            p.paragraph_format.keep_with_next=True
            doc.paragraphs[i+1].paragraph_format.keep_with_next=True


def add_reference_links(doc,references):
    p=style.add_paragraph(doc,'','reference')
    for i,(label,url) in enumerate(references):
        if i:style.format_run(p.add_run('；'),'reference')
        link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(url,RT.HYPERLINK,is_external=True))
        run=style.format_run(p.add_run(label),'reference',bold=True,size_pt=9);p._p.remove(run._r);link.append(run._r);p._p.append(link)


def text_replace(text):
    return REPLACEMENTS.get(text,text)


def add_model_section(doc):
    style.add_paragraph(doc,'小米模型分工','heading2')
    style.add_paragraph(doc,'全部 AI 能力统一由以下五个小米 MiMo 模型提供。文本与视觉场景共用主模型，语音场景分别选择对应模型；常规文字问答不额外调用语音模型。')
    table=style.add_table(doc,['模型','本项目职责','调用与核对要求'],MODELS,widths_cm=[5.1,6.4,6.3])
    for row in table.rows:
        props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
    style.add_paragraph(doc,'图片识别保留人工核对；音频转写不当作身份认证或正式通知原件。声音复刻仅用于本人或已获授权的声音。语音合成不承担推送送达，个人日历与提醒渠道按原方案验证。','note')


def add_integration(doc):
    style.add_paragraph(doc,'小米 AI 对接与验证','heading1')
    style.add_paragraph(doc,'配置与接口','heading2')
    style.add_paragraph(doc,'配置文件位于“系统_开发/.env”。填写 MIMO_API_KEY 后重启，MIMO_BASE_URL 已设为 https://api.xiaomimimo.com/v1。主模型、ASR、普通 TTS、音色设计和声音复刻分别通过 MIMO_TEXT_MODEL、MIMO_ASR_MODEL、MIMO_TTS_MODEL、MIMO_VOICE_DESIGN_MODEL、MIMO_VOICE_CLONE_MODEL 预置为本方案列出的模型。旧供应商配置不再读取。')
    style.add_paragraph(doc,'五个模型统一调用 /chat/completions。主模型沿用 JSON 输出与图片输入，思考开关使用 thinking.type=disabled。ASR 使用 input_audio；TTS 将播报正文置于 assistant 消息，音色描述置于 user 消息，并从返回的 audio.data 读取音频。音色设计不传预置 voice，关闭 optimize_text_preview；声音复刻在 audio.voice 中传入参考录音。')
    style.add_paragraph(doc,'本版本使用小米按量 API。Token Plan 的专用密钥、端点和用途条件需要另行核对，不能直接混用。一个密钥能否调用全部模型，以账号实际权限和额度为准；不把限时优惠或模型配置写成永久免费、已经开通或已经验收。')
    style.add_paragraph(doc,'语音入口与应用上限','heading2')
    style.add_paragraph(doc,'聊天及通知输入框的麦克风调用 ASR。浏览器本地将录音转换为 16 位、16 kHz、单声道 WAV 后发送；录音最长 60 秒。个人设置中的“语音播报”可选择三种 TTS 模式，播报正文最多 2,000 字，音色描述最多 500 字，复刻参考 WAV 最长 30 秒，音频输入不超过 5 MB。以上为本应用限制。')
    style.add_paragraph(doc,'官方语音输入支持 MP3 与 WAV；当前应用先实现 PCM WAV。三种 TTS 首轮统一使用非流式 WAV 输出，生成后由用户点击播放；不把设计、复刻的兼容流式接口当作已经验证的低延迟实时播放。语音数据在应用内存中处理，不保存到通知库；外发前确认本次处理许可。')
    style.add_paragraph(doc,'验收与后续对接','heading2')
    style.add_paragraph(doc,'本轮完成模型配置及请求适配，不代表真实 API 已连通。获得密钥后逐一验证：文字通知提取与引用、截图理解、流式回答、ASR 日期和人名转写、普通播报、音色设计、已授权声音复刻；并核对实际权限、费用、延迟及目标设备录音播放兼容性。声音模型按需要单独验收，不增加日常文本任务的无关调用。')
    style.add_paragraph(doc,'工程接入字段、路由、限制和核验清单详见“系统_开发/MIMO_INTEGRATION.md”。原有本地解析、数据库检索、日期校验、权限、个人确认及日历导出继续由确定性程序处理。全部原件保留，本文件作为小米模型版使用。','note')
    style.add_paragraph(doc,'小米官方接口依据（2026年10月14日核对）','heading2')
    references=[('聊天接口','https://mimo.mi.com/docs/en-US/api/chat'),('结构化输出','https://mimo.mi.com/docs/en-US/quick-start/usage-guide/text-generation/structured-output'),('图片理解','https://mimo.mi.com/docs/en-US/quick-start/usage-guide/multimodal-understanding/image-understanding'),('ASR 接口','https://mimo.mi.com/docs/en-US/api/audio/Speech-Recognition'),('TTS 接口','https://mimo.mi.com/docs/en-US/api/audio/tts')]
    add_reference_links(doc,references)


def convert(path):
    original_hash=hashlib.sha256(path.read_bytes()).hexdigest()
    source=Document(path);doc=style.new_document();added=False
    for element in source.element.body:
        if element.tag==qn('w:p'):
            old=Paragraph(element,source);style_name=old.style.name
            role={'Title':'title','Heading 1':'heading1','Heading 2':'heading2','Caption':'caption','Subtitle':'lead'}.get(style_name,'body')
            if old._p.xpath('.//w:drawing'):
                new=copy.deepcopy(element)
                for blip in new.xpath('.//a:blip'):
                    rid=blip.get(qn('r:embed'))
                    if rid:
                        newrid,_=doc.part.get_or_add_image(io.BytesIO(source.part.related_parts[rid].blob));blip.set(qn('r:embed'),newrid)
                doc.element.body.insert(-1,new);p=Paragraph(new,doc);style.format_paragraph(p,'caption')
                p.paragraph_format.keep_with_next=True
            else:
                text=text_replace(old.text)
                if role=='title':text+='\n'+('最小可行版本' if '最小可行版本' in path.name else '项目方案')+'（小米模型版）'
                p=style.add_paragraph(doc,text,role)
                if role in {'heading1','heading2'}:p.paragraph_format.keep_with_next=True
                if role=='title':style.add_paragraph(doc,'模型配置更新：2026年10月14日','lead')
            if ('首版采用单一模型' in old.text or '协调智能体组织任务' in old.text) and not added:
                add_model_section(doc);added=True
        elif element.tag==qn('w:tbl'):
            old=Table(element,source);rows=[[text_replace(c.text) for c in r.cells] for r in old.rows]
            columns=len(rows[0]);raw_width=[c.width.cm if c.width else 1 for c in old.columns]
            total=sum(raw_width);widths=[v/total*17.8 for v in raw_width]
            table=style.add_table(doc,rows[0],rows[1:],widths_cm=widths)
            for row in table.rows:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    assert added
    add_integration(doc)
    finish_layout(doc)
    out=path.with_name(path.stem+'_小米模型版_2026-10-14.docx')
    style.save(doc,out)
    assert hashlib.sha256(path.read_bytes()).hexdigest()==original_hash
    return {'source':str(path),'source_sha256':original_hash,'output':str(out),'models':[r[0] for r in MODELS]}


if __name__=='__main__':
    records=[convert(FOLDER/name) for name in ['校园通知智能转办助手_项目方案.docx','校园通知智能转办助手_最小可行版本.docx']]
    target=ROOT/'系统_开发/artifacts/mimo-documents.json';target.write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(records,ensure_ascii=False,indent=2))
